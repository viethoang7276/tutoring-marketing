#!/usr/bin/env python3
"""
Render a CMO weekly-report markdown file into a styled Word (.docx) document.

Reusable, pragmatic Markdown -> docx converter scoped to the constructs the reports use:
headings (#/##/###), bullets, checkboxes (- [ ] / - [x]), tables (| ... |), blockquotes (>),
and inline **bold**. Unknown lines fall back to plain paragraphs.

Usage:
    python3 scripts/render_report_docx.py docs/reports/2026-W23-cmo-weekly.md [out.docx]
Importable:
    from render_report_docx import render
    render("docs/reports/2026-W23-cmo-weekly.md", "out.docx")
"""
import sys, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE    = RGBColor(0x46, 0x7D, 0xE3)
BLUE_DK = RGBColor(0x36, 0x69, 0xCC)
INK     = RGBColor(0x02, 0x08, 0x17)
MUTED   = RGBColor(0x64, 0x74, 0x8B)
SANS    = "Calibri"


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _add_inline(paragraph, text, base_color=INK, size=10.5):
    """Render **bold** segments; strip backticks/markdown links to plain text."""
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for i, seg in enumerate(text.split("**")):
        if not seg:
            continue
        r = paragraph.add_run(seg)
        r.font.name = SANS
        r.font.size = Pt(size)
        r.font.color.rgb = base_color
        r.font.bold = (i % 2 == 1)


def _heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = SANS; r.font.bold = True
    if level <= 1:
        r.font.size = Pt(14); r.font.color.rgb = BLUE
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "2"); b.set(qn("w:color"), "467DE3")
        pbdr.append(b); pPr.append(pbdr)
    elif level == 2:
        r.font.size = Pt(12); r.font.color.rgb = INK
    else:
        r.font.size = Pt(10.5); r.font.color.rgb = BLUE_DK


def _table(doc, rows):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols); t.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            val = row[ci] if ci < len(row) else ""
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            _add_inline(p, val, base_color=(RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else INK), size=9)
            if ri == 0:
                for run in p.runs:
                    run.font.bold = True
                _shade(cells[ci], "467DE3")
            elif ri % 2 == 0:
                _shade(cells[ci], "F1F5F9")


def render(md_path, out_path=None):
    md_path = Path(md_path)
    lines = md_path.read_text(encoding="utf8").splitlines()
    out_path = Path(out_path) if out_path else md_path.with_suffix(".docx")

    doc = Document()
    n = doc.styles["Normal"]; n.font.name = SANS; n.font.size = Pt(10.5); n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(5); n.paragraph_format.line_spacing = 1.12

    base = md_path.parent
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        # image:  ![alt](path)  -> embed, centred
        mimg = re.match(r"^!\[[^\]]*\]\(([^)]+)\)\s*$", line.strip())
        if mimg:
            img = (base / mimg.group(1)).resolve()
            if img.exists():
                p = doc.add_paragraph(); p.alignment = 1  # center
                p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
                p.add_run().add_picture(str(img), width=Inches(6.3))
            i += 1; continue
        # table block
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i]); i += 1
            rows = []
            for bl in block:
                if re.match(r"^\s*\|[\s:|-]+\|\s*$", bl):  # separator row
                    continue
                cells = [c.strip() for c in bl.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                _table(doc, rows)
            continue
        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            if level == 1:
                p = doc.add_paragraph(); r = p.add_run(m.group(2))
                r.font.name = SANS; r.font.bold = True; r.font.size = Pt(20); r.font.color.rgb = INK
                p.paragraph_format.space_after = Pt(4)
            else:
                _heading(doc, m.group(2), level - 1)
            i += 1; continue
        # checkbox
        mc = re.match(r"^\s*-\s+\[( |x|X)\]\s+(.*)$", line)
        if mc:
            checked = mc.group(1).lower() == "x"
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            box = p.add_run(("☑  " if checked else "☐  "))
            box.font.name = SANS; box.font.size = Pt(10.5); box.font.color.rgb = BLUE_DK
            _add_inline(p, mc.group(2))
            i += 1; continue
        # bullet
        mb = re.match(r"^\s*[-*]\s+(.*)$", line)
        if mb:
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
            _add_inline(p, mb.group(1))
            i += 1; continue
        # blockquote
        if line.lstrip().startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            _add_inline(p, line.lstrip().lstrip(">").strip(), base_color=MUTED, size=10)
            for r in p.runs:
                r.font.italic = True
            i += 1; continue
        # whole-line italic caption:  _text_
        mi = re.match(r"^_(.+)_$", line.strip())
        if mi:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            r = p.add_run(mi.group(1))
            r.font.italic = True; r.font.name = SANS; r.font.size = Pt(9.5); r.font.color.rgb = MUTED
            i += 1; continue
        # plain paragraph — join consecutive soft-wrapped lines into ONE paragraph
        # (Markdown rule: a blank line separates paragraphs; single newlines are soft wraps).
        buf = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            stripped = nxt.lstrip()
            if (not nxt.strip()
                    or stripped[:1] in "|#>"
                    or stripped.startswith("![")
                    or re.match(r"^[-*]\s", stripped)):
                break
            buf.append(stripped)
            i += 1
        p = doc.add_paragraph(); _add_inline(p, " ".join(buf))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return str(out_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: render_report_docx.py <report.md> [out.docx]"); sys.exit(1)
    out = render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print("Rendered:", out)
